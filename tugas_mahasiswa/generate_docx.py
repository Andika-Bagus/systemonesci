import os
import re
import sys
import subprocess

# Ensure python-docx is installed
try:
    import docx
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
except ImportError:
    print("python-docx not found. Installing now...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding of a table cell (values in twips: 20 twips = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table):
    """Apply thin light gray borders to the table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_callout_box(paragraph, text, fill_hex="F2F5F8", border_hex="1F4E79"):
    """Format a paragraph as a callout box with left border and light shading."""
    pPr = paragraph._p.get_or_add_pPr()
    
    # Background shading
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    pPr.append(shd)
    
    # Left border
    pbdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')  # 3pt thickness
    left.set(qn('w:space'), '12') # Padding space
    left.set(qn('w:color'), border_hex)
    pbdr.append(left)
    pPr.append(pbdr)
    
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.25)

def main():
    doc = docx.Document()
    
    # Page setup
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    
    # Define styles & colors
    COLOR_NAVY = RGBColor(31, 78, 121)    # #1F4E79
    COLOR_CHARCOAL = RGBColor(51, 51, 51) # #333333
    COLOR_GRAY = RGBColor(112, 128, 144)  # #708090
    
    # ------------------ COVER PAGE ------------------
    for _ in range(5):
        doc.add_paragraph()
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("PAKET DOKUMEN TUGAS MAHASISWA")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_NAVY
    
    p_subtitle = doc.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_subtitle.add_run(
        "Sistem Pemantauan Performa Website (PageSpeed Monitor) &\n"
        "Manajemen Otentikasi Sesi Aman Kredensial Jurnal (OJS Secure)"
    )
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(14)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_GRAY
    
    for _ in range(8):
        doc.add_paragraph()
        
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run(
        "Diajukan untuk Memenuhi Tugas Latihan Rekayasa Perangkat Lunak / Magang\n\n"
        "Lokasi Magang: CV Syntax Corporation Indonesia (Syntax Indonesia)\n"
        "Tahun Akademik: 2026\n"
    )
    run_meta.font.name = 'Arial'
    run_meta.font.size = Pt(11)
    run_meta.font.color.rgb = COLOR_CHARCOAL
    
    p_dev = doc.add_paragraph()
    p_dev.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_dev = p_dev.add_run("Disusun Oleh:\nTim Mahasiswa Magang IT")
    run_dev.font.name = 'Arial'
    run_dev.font.size = Pt(12)
    run_dev.font.bold = True
    run_dev.font.color.rgb = COLOR_NAVY
    
    doc.add_page_break()
    
    # ------------------ DEFINE FILES TO MERGE ------------------
    files = [
        ("1. Draft Kesepakatan Proyek", "1_draft_kesepakatan_proyek.md"),
        ("2. Business Requirements Document (BRD)", "2_brd_business_requirements_document.md"),
        ("3. Rumusan Masalah Penelitian", "3_rumusan_masalah_penelitian.md"),
        ("4. Rancangan Metode DSR", "4_rancangan_metode_dsr.md"),
        ("5. Rencana Pengujian & Instrumen", "5_rencana_pengujian_instrumen.md"),
    ]
    
    base_path = os.path.dirname(os.path.abspath(__file__))
    
    for title_sec, filename in files:
        file_path = os.path.join(base_path, filename)
        if not os.path.exists(file_path):
            print(f"Warning: File {filename} not found, skipping...")
            continue
            
        print(f"Parsing and appending {filename}...")
        
        # Section Header Page Break
        p_sec_title = doc.add_paragraph()
        p_sec_title.paragraph_format.space_before = Pt(18)
        p_sec_title.paragraph_format.space_after = Pt(12)
        p_sec_title.paragraph_format.keep_with_next = True
        run_sec_title = p_sec_title.add_run(title_sec.upper())
        run_sec_title.font.name = 'Arial'
        run_sec_title.font.size = Pt(18)
        run_sec_title.font.bold = True
        run_sec_title.font.color.rgb = COLOR_NAVY
        
        # Parse the Markdown file content
        with open(file_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()
            
        in_table = False
        table_rows = []
        in_code = False
        code_lines = []
        
        for line in lines:
            stripped = line.strip()
            
            # Skip main headers since we add them manually above
            if stripped.startswith("# ") and not (stripped.startswith("## ") or stripped.startswith("### ")):
                continue
                
            # Handle Code Block (or Mermaid)
            if stripped.startswith("```"):
                if in_code:
                    # End code block and add callout box
                    in_code = False
                    p_code = doc.add_paragraph()
                    p_code.paragraph_format.space_before = Pt(6)
                    p_code.paragraph_format.space_after = Pt(6)
                    code_text = "\n".join(code_lines)
                    run_code = p_code.add_run(code_text)
                    run_code.font.name = 'Consolas'
                    run_code.font.size = Pt(9.5)
                    run_code.font.color.rgb = RGBColor(60, 76, 92)
                    make_callout_box(p_code, code_text, fill_hex="F4F6F8", border_hex="708090")
                    code_lines = []
                else:
                    in_code = True
                continue
                
            if in_code:
                code_lines.append(line.rstrip('\n'))
                continue
                
            # Handle Table parsing
            if stripped.startswith("|"):
                # Detect separator line e.g. |---|---|
                if re.match(r'^\|[\s:-|]*\|$', stripped):
                    continue
                in_table = True
                # Parse columns
                cols = [c.strip() for c in stripped.split("|")[1:-1]]
                table_rows.append(cols)
                continue
            else:
                if in_table:
                    # End of table, compile and add table to word
                    if table_rows:
                        num_cols = len(table_rows[0])
                        w_table = doc.add_table(rows=len(table_rows), cols=num_cols)
                        w_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                        set_table_borders(w_table)
                        
                        for row_idx, row_data in enumerate(table_rows):
                            w_row = w_table.rows[row_idx]
                            is_header = (row_idx == 0)
                            
                            for col_idx, col_val in enumerate(row_data):
                                if col_idx >= num_cols:
                                    break
                                cell = w_row.cells[col_idx]
                                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                                set_cell_margins(cell, top=140, bottom=140, left=180, right=180) # Twips padding
                                
                                # Set background shading
                                if is_header:
                                    set_cell_background(cell, "1F4E79") # Navy header
                                elif row_idx % 2 == 0:
                                    set_cell_background(cell, "F2F5F8") # Zebra striping
                                else:
                                    set_cell_background(cell, "FFFFFF")
                                    
                                # Cell text
                                p_cell = cell.paragraphs[0]
                                p_cell.paragraph_format.space_before = Pt(2)
                                p_cell.paragraph_format.space_after = Pt(2)
                                
                                # Text styling
                                run_cell = p_cell.add_run(col_val)
                                run_cell.font.name = 'Arial'
                                run_cell.font.size = Pt(10)
                                if is_header:
                                    run_cell.font.bold = True
                                    run_cell.font.color.rgb = RGBColor(255, 255, 255) # White text
                                else:
                                    run_cell.font.color.rgb = COLOR_CHARCOAL
                                    
                        doc.add_paragraph() # spacing after table
                    in_table = False
                    table_rows = []
                    
            # Skip empty lines
            if not stripped:
                continue
                
            # Handle Headers (H2 & H3)
            if stripped.startswith("## "):
                h_text = stripped[3:].strip()
                p_h = doc.add_paragraph()
                p_h.paragraph_format.space_before = Pt(14)
                p_h.paragraph_format.space_after = Pt(6)
                p_h.paragraph_format.keep_with_next = True
                run_h = p_h.add_run(h_text)
                run_h.font.name = 'Arial'
                run_h.font.size = Pt(14)
                run_h.font.bold = True
                run_h.font.color.rgb = COLOR_NAVY
                
            elif stripped.startswith("### "):
                h_text = stripped[4:].strip()
                p_h = doc.add_paragraph()
                p_h.paragraph_format.space_before = Pt(10)
                p_h.paragraph_format.space_after = Pt(4)
                p_h.paragraph_format.keep_with_next = True
                run_h = p_h.add_run(h_text)
                run_h.font.name = 'Arial'
                run_h.font.size = Pt(12)
                run_h.font.bold = True
                run_h.font.color.rgb = COLOR_NAVY
                
            # Handle Bullet points / Lists
            elif stripped.startswith("* ") or stripped.startswith("- "):
                list_text = stripped[2:].strip()
                p_li = doc.add_paragraph(style='List Bullet')
                p_li.paragraph_format.space_before = Pt(2)
                p_li.paragraph_format.space_after = Pt(2)
                run_li = p_li.add_run(list_text)
                run_li.font.name = 'Arial'
                run_li.font.size = Pt(11)
                run_li.font.color.rgb = COLOR_CHARCOAL
                
            # Handle Numbered Lists e.g. "1. " or "2. "
            elif re.match(r'^\d+\.\s', stripped):
                list_text = stripped[stripped.find(".")+1:].strip()
                p_li = doc.add_paragraph(style='List Number')
                p_li.paragraph_format.space_before = Pt(2)
                p_li.paragraph_format.space_after = Pt(2)
                run_li = p_li.add_run(list_text)
                run_li.font.name = 'Arial'
                run_li.font.size = Pt(11)
                run_li.font.color.rgb = COLOR_CHARCOAL
                
            # Handle Regular Paragraphs
            else:
                # Basic inline markdown bold/italic rendering
                p_para = doc.add_paragraph()
                p_para.paragraph_format.space_before = Pt(4)
                p_para.paragraph_format.space_after = Pt(6)
                p_para.paragraph_format.line_spacing = 1.15
                
                # Simple parser for bold (**text**)
                tokens = re.split(r'(\*\*.*?\*\*)', stripped)
                for token in tokens:
                    if token.startswith("**") and token.endswith("**"):
                        run_token = p_para.add_run(token[2:-2])
                        run_token.font.bold = True
                    else:
                        # Simple parser for italic (*text*)
                        sub_tokens = re.split(r'(\*.*?\*)', token)
                        for sub_token in sub_tokens:
                            if sub_token.startswith("*") and sub_token.endswith("*"):
                                run_sub = p_para.add_run(sub_token[1:-1])
                                run_sub.font.italic = True
                            else:
                                run_sub = p_para.add_run(sub_token)
                                
                    for run in p_para.runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(11)
                        run.font.color.rgb = COLOR_CHARCOAL
                        
        doc.add_page_break()
        
    # Save the output file
    output_path = os.path.join(base_path, "Tugas_Mahasiswa_Proyek_OJS_Secure.docx")
    doc.save(output_path)
    print(f"Successfully compiled all documents into: {output_path}")

if __name__ == "__main__":
    main()
